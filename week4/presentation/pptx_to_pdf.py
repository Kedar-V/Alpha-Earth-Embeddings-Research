#!/usr/bin/env python3
"""Rasterize a .pptx to PDF without LibreOffice (pictures, text, tables, fills)."""
from __future__ import annotations

import argparse
import io
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu


def emu_to_px(emu: int, dpi: int) -> int:
    return int(round(emu * dpi / 914400))


def pt_to_px(pt: float, dpi: int) -> int:
    return max(8, int(round(pt * dpi / 72)))


def rgb_from_color(color) -> tuple[int, int, int] | None:
    if color is None:
        return None
    try:
        if color.type is not None and str(color.type).endswith("RGB"):
            rgb = color.rgb
            if isinstance(rgb, RGBColor):
                return (rgb[0], rgb[1], rgb[2])
    except Exception:
        pass
    return None


def load_font(size_px: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    ]
    for path in candidates:
        if Path(path).is_file():
            return ImageFont.truetype(path, size=size_px)
    return ImageFont.load_default()


def wrap_text(text: str, font, max_width: int, draw: ImageDraw.ImageDraw) -> list[str]:
    words = text.replace("\r", "").split()
    if not words:
        return [""]
    lines: list[str] = []
    cur = words[0]
    for w in words[1:]:
        trial = f"{cur} {w}"
        if draw.textlength(trial, font=font) <= max_width:
            cur = trial
        else:
            lines.append(cur)
            cur = w
    lines.append(cur)
    return lines


def draw_text_frame(draw, shape, dpi: int):
    if not shape.has_text_frame:
        return
    left = emu_to_px(shape.left, dpi)
    top = emu_to_px(shape.top, dpi)
    width = max(1, emu_to_px(shape.width, dpi))
    height = max(1, emu_to_px(shape.height, dpi))
    y = top
    for para in shape.text_frame.paragraphs:
        text = para.text.strip()
        if not text:
            y += pt_to_px(6, dpi)
            continue
        size_pt = 14.0
        bold = False
        color = (10, 25, 47)
        if para.runs:
            r = para.runs[0]
            if r.font.size is not None:
                size_pt = r.font.size.pt
            if r.font.bold:
                bold = True
            c = rgb_from_color(r.font.color)
            if c:
                color = c
        font = load_font(pt_to_px(size_pt, dpi), bold=bold)
        for line in wrap_text(text, font, width - 4, draw):
            if y > top + height:
                break
            draw.text((left + 2, y), line, fill=color, font=font)
            y += pt_to_px(size_pt * 1.25, dpi)


def draw_table(draw, shape, dpi: int):
    table = shape.table
    left = emu_to_px(shape.left, dpi)
    top = emu_to_px(shape.top, dpi)
    width = emu_to_px(shape.width, dpi)
    height = emu_to_px(shape.height, dpi)
    rows, cols = len(table.rows), len(table.columns)
    if rows == 0 or cols == 0:
        return
    cell_w = width // cols
    cell_h = height // rows
    draw.rectangle([left, top, left + width, top + height], outline=(30, 144, 202), width=1)
    for r in range(rows):
        for c in range(cols):
            cx = left + c * cell_w
            cy = top + r * cell_h
            draw.line([(cx, top), (cx, top + height)], fill=(200, 210, 220), width=1)
            draw.line([(left, cy), (left + width, cy)], fill=(200, 210, 220), width=1)
            cell = table.cell(r, c)
            text = cell.text.strip()
            if not text:
                continue
            size_pt = 12.0 if r > 0 else 11.0
            bold = r == 0
            font = load_font(pt_to_px(size_pt, dpi), bold=bold)
            color = (30, 144, 202) if r == 0 else (10, 25, 47)
            pad = 6
            for i, line in enumerate(wrap_text(text, font, cell_w - 2 * pad, draw)):
                draw.text((cx + pad, cy + pad + i * pt_to_px(size_pt * 1.2, dpi)), line, fill=color, font=font)


def render_slide(slide, size: tuple[int, int], dpi: int) -> Image.Image:
    img = Image.new("RGB", size, (255, 255, 255))
    draw = ImageDraw.Draw(img)
    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            blob = shape.image.blob
            pic = Image.open(io.BytesIO(blob)).convert("RGBA")
            left = emu_to_px(shape.left, dpi)
            top = emu_to_px(shape.top, dpi)
            width = max(1, emu_to_px(shape.width, dpi))
            height = max(1, emu_to_px(shape.height, dpi))
            pic = pic.resize((width, height), Image.Resampling.LANCZOS)
            img.paste(pic, (left, top), pic)
        elif shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            fill = None
            try:
                fill = rgb_from_color(shape.fill.fore_color)
            except Exception:
                fill = None
            if fill:
                left = emu_to_px(shape.left, dpi)
                top = emu_to_px(shape.top, dpi)
                right = left + emu_to_px(shape.width, dpi)
                bottom = top + emu_to_px(shape.height, dpi)
                draw.rectangle([left, top, right, bottom], fill=fill)
            draw_text_frame(draw, shape, dpi)
        elif shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
            draw_text_frame(draw, shape, dpi)
        elif shape.shape_type == MSO_SHAPE_TYPE.TABLE:
            draw_table(draw, shape, dpi)
    return img


def pptx_to_pdf(pptx_path: Path, pdf_path: Path, dpi: int = 150) -> None:
    prs = Presentation(str(pptx_path))
    w = emu_to_px(prs.slide_width, dpi)
    h = emu_to_px(prs.slide_height, dpi)
    pages: list[Image.Image] = []
    for slide in prs.slides:
        pages.append(render_slide(slide, (w, h), dpi))
    if not pages:
        raise SystemExit("empty deck")
    first, rest = pages[0], pages[1:]
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    first.save(
        pdf_path,
        "PDF",
        resolution=dpi,
        save_all=True,
        append_images=rest,
    )


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("-o", "--out", required=True)
    ap.add_argument("--dpi", type=int, default=150)
    args = ap.parse_args(argv)
    pptx_to_pdf(Path(args.pptx), Path(args.out), dpi=args.dpi)
    print(f"wrote {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
